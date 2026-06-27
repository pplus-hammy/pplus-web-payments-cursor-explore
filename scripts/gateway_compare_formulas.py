#!/usr/bin/env python3
"""Gateway Performance Comparison — Formula-based Excel Output

Reads the same xlsx input as gateway_compare.py but produces an Excel workbook
where every derived value is an Excel formula referencing source cells. This
makes the spreadsheet fully auditable and allows stakeholders to tweak
assumptions (cost rates, lifetime, avg txn) and see results recalculate live.

Usage:
    1. Set INPUT_DIR, INPUT_FILENAME, and OUTPUT_DIR at the top of this script
    2. python3 scripts/gateway_compare_formulas.py

Output:
    gateway_comparison_formulas.xlsx in the output directory, with sheets:
      1. Parameters        — editable source values
      2. Raw Data          — input data from the xlsx (values only)
      3. Monthly Comparison — formulas referencing Raw Data + Parameters
      4. CLV Lost to Failures — formulas for failure-cost analysis
      5. What-If Analysis  — all volume through one gateway (formulas)
      6. Volume-Shift      — 10%/20%/25% DTP split scenarios (formulas)
"""

import sys
import os
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter

GW_A = "Adyen"
GW_C = "Cybersource"

# ── Input/Output Configuration ───────────────────────────────────────────
INPUT_DIR = "/Users/gregory.hamilton/Downloads"
INPUT_FILENAME = "gateway_performance_test.xlsx"
OUTPUT_DIR = os.path.join(INPUT_DIR, "gateway_comparison_output")

NS_STRICT = "http://purl.oclc.org/ooxml/spreadsheetml/main"
NS_TRANS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


# ═══════════════════════════════════════════════════════════════════════════
#  DATA LOADING  (shared with gateway_compare.py)
# ═══════════════════════════════════════════════════════════════════════════

def _detect_ns(xml_bytes):
    snippet = xml_bytes[:2000].decode("utf-8", errors="ignore")
    return NS_STRICT if NS_STRICT in snippet else NS_TRANS


def _parse_shared_strings(xml_bytes, ns):
    root = ET.fromstring(xml_bytes)
    return [
        (si.find(f"{{{ns}}}t").text or "")
        if si.find(f"{{{ns}}}t") is not None
        else ""
        for si in root.findall(f"{{{ns}}}si")
    ]


def _col_to_idx(col):
    idx = 0
    for ch in col:
        idx = idx * 26 + (ord(ch.upper()) - 64)
    return idx - 1


def _parse_sheet(xml_bytes, ns, shared):
    root = ET.fromstring(xml_bytes)
    rows = {}
    for row_el in root.iter(f"{{{ns}}}row"):
        rn = int(row_el.get("r"))
        cells = {}
        for c in row_el.findall(f"{{{ns}}}c"):
            ref = c.get("r")
            col = _col_to_idx("".join(ch for ch in ref if ch.isalpha()))
            v = c.find(f"{{{ns}}}v")
            if v is None or v.text is None:
                continue
            raw = v.text
            ct = c.get("t", "")
            if ct == "s":
                cells[col] = shared[int(raw)]
            elif ct == "d":
                cells[col] = raw
            else:
                try:
                    cells[col] = int(raw) if "." not in raw else float(raw)
                except ValueError:
                    cells[col] = raw
        rows[rn] = cells
    return rows


def load_data(filepath):
    """Return (params, month_list) parsed from the xlsx."""
    with zipfile.ZipFile(filepath) as zf:
        ss_xml = zf.read("xl/sharedStrings.xml")
        ns = _detect_ns(ss_xml)
        shared = _parse_shared_strings(ss_xml, ns)
        sh_xml = zf.read("xl/worksheets/sheet1.xml")
        rows = _parse_sheet(sh_xml, _detect_ns(sh_xml), shared)

    params = {
        "cost_a": float(rows.get(1, {}).get(1, 0)),
        "cost_c": float(rows.get(2, {}).get(1, 0)),
        "avg_txn": float(rows.get(3, {}).get(1, 0)),
        "lifetime": int(rows.get(4, {}).get(1, 0)),
    }

    months = []
    for rn in sorted(rows):
        if rn < 7:
            continue
        r = rows[rn]
        if 0 not in r:
            continue
        ds = r[0]
        if isinstance(ds, str):
            if "-" in ds and len(ds) == 10:
                try:
                    dt = datetime.strptime(ds, "%Y-%m-%d")
                    label = dt.strftime("%b %Y")
                except ValueError:
                    label = ds
            else:
                label = str(ds)
        elif isinstance(ds, (int, float)):
            label = str(ds)
        else:
            label = str(ds)

        months.append({
            "label": label,
            "sr_ci_a": float(r.get(1, 0)), "sr_ci_c": float(r.get(2, 0)),
            "sr_ren_a": float(r.get(3, 0)), "sr_ren_c": float(r.get(4, 0)),
            "vol_ci_a": int(r.get(5, 0)), "vol_ci_c": int(r.get(6, 0)),
            "vol_ren_a": int(r.get(7, 0)), "vol_ren_c": int(r.get(8, 0)),
            "ref_ci_a": float(r.get(9, 0)), "ref_ci_c": float(r.get(10, 0)),
            "ref_ren_a": float(r.get(11, 0)), "ref_ren_c": float(r.get(12, 0)),
            "cb_ci_a": float(r.get(13, 0)), "cb_ci_c": float(r.get(14, 0)),
            "cb_ren_a": float(r.get(15, 0)), "cb_ren_c": float(r.get(16, 0)),
        })
    return params, months


# ═══════════════════════════════════════════════════════════════════════════
#  STYLING HELPERS
# ═══════════════════════════════════════════════════════════════════════════

HDR_FONT = Font(bold=True, size=11)
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTALS_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin")
HDR_BORDER = Border(bottom=THIN)

MONEY = '#,##0.00'
PCT = '0.00%'
INT = '#,##0'


def _style_header(ws, row_num, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = HDR_BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _style_totals_row(ws, row_num, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HDR_FONT
        cell.fill = TOTALS_FILL


def _set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _cell(ws, row, col, value=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if fmt:
        c.number_format = fmt
    return c


# ═══════════════════════════════════════════════════════════════════════════
#  PARAMETER CELL REFERENCES (absolute, for cross-sheet formulas)
# ═══════════════════════════════════════════════════════════════════════════

P_COST_A = "Parameters!$B$2"
P_COST_C = "Parameters!$B$3"
P_AVG_TXN = "Parameters!$B$4"
P_LIFETIME = "Parameters!$B$5"
P_CLV = "Parameters!$B$6"
P_FUTURE_REN = "Parameters!$B$7"


# ═══════════════════════════════════════════════════════════════════════════
#  RAW DATA CELL REFERENCE HELPERS
# ═══════════════════════════════════════════════════════════════════════════

# Raw Data columns (1-indexed):
# A=Month  B=SR_DTP_A  C=SR_DTP_C  D=SR_Ren_A  E=SR_Ren_C
# F=Vol_DTP_A  G=Vol_DTP_C  H=Vol_Ren_A  I=Vol_Ren_C
# J=Ref_DTP_A  K=Ref_DTP_C  L=Ref_Ren_A  M=Ref_Ren_C
# N=CB_DTP_A  O=CB_DTP_C  P=CB_Ren_A  Q=CB_Ren_C

def _rd(col_letter, row):
    return f"'Raw Data'!{col_letter}{row}"


# ═══════════════════════════════════════════════════════════════════════════
#  SHEET BUILDERS
# ═══════════════════════════════════════════════════════════════════════════

def _build_parameters(wb, params):
    ws = wb.create_sheet("Parameters", 0)
    data = [
        ("Parameter", "Value"),
        (f"{GW_A} Cost Rate", params["cost_a"]),
        (f"{GW_C} Cost Rate", params["cost_c"]),
        ("Average Transaction Amount", params["avg_txn"]),
        ("Customer Lifetime (months)", params["lifetime"]),
        (None, None),
        (None, None),
    ]
    for ri, (label, val) in enumerate(data, 1):
        ws.cell(row=ri, column=1, value=label)
        _cell(ws, ri, 2, val)

    _cell(ws, 1, 1).font = HDR_FONT
    _cell(ws, 1, 2).font = HDR_FONT

    _cell(ws, 2, 2).number_format = PCT
    _cell(ws, 3, 2).number_format = PCT
    _cell(ws, 4, 2).number_format = MONEY
    _cell(ws, 5, 2).number_format = INT

    ws.cell(row=6, column=1, value="CLV per Customer")
    _cell(ws, 6, 2, f"={P_LIFETIME}*{P_AVG_TXN}", MONEY)

    ws.cell(row=7, column=1, value="Future Renewals Lost per Failed DTP")
    _cell(ws, 7, 2, f"=({P_LIFETIME}-1)*{P_AVG_TXN}", MONEY)

    _set_col_widths(ws, [38, 20])
    return ws


def _build_raw_data(wb, month_list):
    ws = wb.create_sheet("Raw Data")
    headers = [
        "Month",
        f"SR DTP ({GW_A})", f"SR DTP ({GW_C})",
        f"SR Ren ({GW_A})", f"SR Ren ({GW_C})",
        f"Vol DTP ({GW_A})", f"Vol DTP ({GW_C})",
        f"Vol Ren ({GW_A})", f"Vol Ren ({GW_C})",
        f"Ref/DTP ({GW_A})", f"Ref/DTP ({GW_C})",
        f"Ref/Ren ({GW_A})", f"Ref/Ren ({GW_C})",
        f"CB/DTP ({GW_A})", f"CB/DTP ({GW_C})",
        f"CB/Ren ({GW_A})", f"CB/Ren ({GW_C})",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))

    KEYS = [
        "label",
        "sr_ci_a", "sr_ci_c", "sr_ren_a", "sr_ren_c",
        "vol_ci_a", "vol_ci_c", "vol_ren_a", "vol_ren_c",
        "ref_ci_a", "ref_ci_c", "ref_ren_a", "ref_ren_c",
        "cb_ci_a", "cb_ci_c", "cb_ren_a", "cb_ren_c",
    ]
    for ri, md in enumerate(month_list, 2):
        for ci, k in enumerate(KEYS, 1):
            val = md[k]
            c = _cell(ws, ri, ci, val)
            if ci in (2, 3, 4, 5):
                c.number_format = PCT
            elif ci in (6, 7, 8, 9):
                c.number_format = INT
            elif ci >= 10:
                c.number_format = MONEY

    _set_col_widths(ws, [14] + [16] * 16)
    return ws, len(month_list)


def _build_monthly_comparison(wb, num_months):
    """Monthly Comparison sheet — all derived values are formulas."""
    ws = wb.create_sheet("Monthly Comparison")

    # Column layout (1-indexed):
    #  A: Month
    #  B: DTP Attempts (A)          C: DTP Attempts (C)
    #  D: DTP SR (A)                E: DTP SR (C)
    #  F: DTP Successful (A)        G: DTP Successful (C)
    #  H: DTP Failed (A)            I: DTP Failed (C)
    #  J: Ren Attempts (A)          K: Ren Attempts (C)
    #  L: Ren SR (A)                M: Ren SR (C)
    #  N: Ren Successful (A)        O: Ren Successful (C)
    #  P: Ren Failed (A)            Q: Ren Failed (C)
    #  R: Total Attempts (A)        S: Total Attempts (C)
    #  T: Total Successful (A)      U: Total Successful (C)
    #  V: Overall SR (A)            W: Overall SR (C)
    #  X: Gross Revenue (A)         Y: Gross Revenue (C)
    #  Z: Gateway Cost (A)          AA: Gateway Cost (C)
    #  AB: Refunds (A)              AC: Refunds (C)
    #  AD: Chargebacks (A)          AE: Chargebacks (C)
    #  AF: Net Revenue (A)          AG: Net Revenue (C)
    #  AH: Net/Attempt (A)          AI: Net/Attempt (C)
    #  AJ: Net/Success (A)          AK: Net/Success (C)

    headers = [
        "Month",
        f"DTP Attempts ({GW_A})", f"DTP Attempts ({GW_C})",
        f"DTP SR ({GW_A})", f"DTP SR ({GW_C})",
        f"DTP Successful ({GW_A})", f"DTP Successful ({GW_C})",
        f"DTP Failed ({GW_A})", f"DTP Failed ({GW_C})",
        f"Ren Attempts ({GW_A})", f"Ren Attempts ({GW_C})",
        f"Ren SR ({GW_A})", f"Ren SR ({GW_C})",
        f"Ren Successful ({GW_A})", f"Ren Successful ({GW_C})",
        f"Ren Failed ({GW_A})", f"Ren Failed ({GW_C})",
        f"Total Attempts ({GW_A})", f"Total Attempts ({GW_C})",
        f"Total Successful ({GW_A})", f"Total Successful ({GW_C})",
        f"Overall SR ({GW_A})", f"Overall SR ({GW_C})",
        f"Gross Revenue ({GW_A})", f"Gross Revenue ({GW_C})",
        f"Gateway Cost ({GW_A})", f"Gateway Cost ({GW_C})",
        f"Refunds ({GW_A})", f"Refunds ({GW_C})",
        f"Chargebacks ({GW_A})", f"Chargebacks ({GW_C})",
        f"Net Revenue ({GW_A})", f"Net Revenue ({GW_C})",
        f"Net/Attempt ({GW_A})", f"Net/Attempt ({GW_C})",
        f"Net/Success ({GW_A})", f"Net/Success ({GW_C})",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))

    first_data = 2
    last_data = first_data + num_months - 1
    totals_row = last_data + 1

    for r in range(first_data, last_data + 1):
        rd_r = r  # Raw Data rows align (both start at row 2)

        # A: Month label
        _cell(ws, r, 1, f"={_rd('A', rd_r)}")

        # B,C: DTP Attempts — direct from raw data
        _cell(ws, r, 2, f"={_rd('F', rd_r)}", INT)
        _cell(ws, r, 3, f"={_rd('G', rd_r)}", INT)

        # D,E: DTP Success Rate — direct from raw data
        _cell(ws, r, 4, f"={_rd('B', rd_r)}", PCT)
        _cell(ws, r, 5, f"={_rd('C', rd_r)}", PCT)

        # F,G: DTP Successful = ROUND(Attempts * SR)
        _cell(ws, r, 6, f"=ROUND(B{r}*D{r},0)", INT)
        _cell(ws, r, 7, f"=ROUND(C{r}*E{r},0)", INT)

        # H,I: DTP Failed = Attempts - Successful
        _cell(ws, r, 8, f"=B{r}-F{r}", INT)
        _cell(ws, r, 9, f"=C{r}-G{r}", INT)

        # J,K: Ren Attempts
        _cell(ws, r, 10, f"={_rd('H', rd_r)}", INT)
        _cell(ws, r, 11, f"={_rd('I', rd_r)}", INT)

        # L,M: Ren Success Rate
        _cell(ws, r, 12, f"={_rd('D', rd_r)}", PCT)
        _cell(ws, r, 13, f"={_rd('E', rd_r)}", PCT)

        # N,O: Ren Successful = ROUND(Attempts * SR)
        _cell(ws, r, 14, f"=ROUND(J{r}*L{r},0)", INT)
        _cell(ws, r, 15, f"=ROUND(K{r}*M{r},0)", INT)

        # P,Q: Ren Failed
        _cell(ws, r, 16, f"=J{r}-N{r}", INT)
        _cell(ws, r, 17, f"=K{r}-O{r}", INT)

        # R,S: Total Attempts = DTP + Ren
        _cell(ws, r, 18, f"=B{r}+J{r}", INT)
        _cell(ws, r, 19, f"=C{r}+K{r}", INT)

        # T,U: Total Successful = DTP Succ + Ren Succ
        _cell(ws, r, 20, f"=F{r}+N{r}", INT)
        _cell(ws, r, 21, f"=G{r}+O{r}", INT)

        # V,W: Overall SR = Total Succ / Total Attempts
        _cell(ws, r, 22, f"=IF(R{r}=0,0,T{r}/R{r})", PCT)
        _cell(ws, r, 23, f"=IF(S{r}=0,0,U{r}/S{r})", PCT)

        # X,Y: Gross Revenue = Total Succ * Avg Txn
        _cell(ws, r, 24, f"=T{r}*{P_AVG_TXN}", MONEY)
        _cell(ws, r, 25, f"=U{r}*{P_AVG_TXN}", MONEY)

        # Z,AA: Gateway Cost = Gross * Cost Rate
        _cell(ws, r, 26, f"=X{r}*{P_COST_A}", MONEY)
        _cell(ws, r, 27, f"=Y{r}*{P_COST_C}", MONEY)

        # AB,AC: Refunds = DTP_Succ*Ref/DTP + Ren_Succ*Ref/Ren
        _cell(ws, r, 28, f"=F{r}*{_rd('J', rd_r)}+N{r}*{_rd('L', rd_r)}", MONEY)
        _cell(ws, r, 29, f"=G{r}*{_rd('K', rd_r)}+O{r}*{_rd('M', rd_r)}", MONEY)

        # AD,AE: Chargebacks = DTP_Succ*CB/DTP + Ren_Succ*CB/Ren
        _cell(ws, r, 30, f"=F{r}*{_rd('N', rd_r)}+N{r}*{_rd('P', rd_r)}", MONEY)
        _cell(ws, r, 31, f"=G{r}*{_rd('O', rd_r)}+O{r}*{_rd('Q', rd_r)}", MONEY)

        # AF,AG: Net Revenue = Gross - Cost - Refunds - Chargebacks
        _cell(ws, r, 32, f"=X{r}-Z{r}-AB{r}-AD{r}", MONEY)
        _cell(ws, r, 33, f"=Y{r}-AA{r}-AC{r}-AE{r}", MONEY)

        # AH,AI: Net per Attempt
        _cell(ws, r, 34, f"=IF(R{r}=0,0,AF{r}/R{r})", MONEY)
        _cell(ws, r, 35, f"=IF(S{r}=0,0,AG{r}/S{r})", MONEY)

        # AJ,AK: Net per Successful Txn
        _cell(ws, r, 36, f"=IF(T{r}=0,0,AF{r}/T{r})", MONEY)
        _cell(ws, r, 37, f"=IF(U{r}=0,0,AG{r}/U{r})", MONEY)

    # Totals row
    tr = totals_row
    rng = f"{first_data}:{last_data}"
    _cell(ws, tr, 1, f"{num_months}-Month Total")

    sum_cols = [2,3, 6,7, 8,9, 10,11, 14,15, 16,17, 18,19, 20,21, 24,25, 26,27, 28,29, 30,31, 32,33]
    for ci in sum_cols:
        cl = get_column_letter(ci)
        fmt = MONEY if ci >= 24 else INT
        _cell(ws, tr, ci, f"=SUM({cl}{first_data}:{cl}{last_data})", fmt)

    # Weighted avg SR for totals: Succ / Attempts
    _cell(ws, tr, 4, f"=IF(B{tr}=0,0,F{tr}/B{tr})", PCT)
    _cell(ws, tr, 5, f"=IF(C{tr}=0,0,G{tr}/C{tr})", PCT)
    _cell(ws, tr, 12, f"=IF(J{tr}=0,0,N{tr}/J{tr})", PCT)
    _cell(ws, tr, 13, f"=IF(K{tr}=0,0,O{tr}/K{tr})", PCT)
    _cell(ws, tr, 22, f"=IF(R{tr}=0,0,T{tr}/R{tr})", PCT)
    _cell(ws, tr, 23, f"=IF(S{tr}=0,0,U{tr}/S{tr})", PCT)

    _cell(ws, tr, 34, f"=IF(R{tr}=0,0,AF{tr}/R{tr})", MONEY)
    _cell(ws, tr, 35, f"=IF(S{tr}=0,0,AG{tr}/S{tr})", MONEY)
    _cell(ws, tr, 36, f"=IF(T{tr}=0,0,AF{tr}/T{tr})", MONEY)
    _cell(ws, tr, 37, f"=IF(U{tr}=0,0,AG{tr}/U{tr})", MONEY)

    _style_totals_row(ws, tr, len(headers))
    _set_col_widths(ws, [14] + [18] * (len(headers) - 1))
    return ws, first_data, last_data, totals_row


def _build_clv_lost(wb, num_months, mc_first, mc_last):
    """CLV Lost to Failures — formulas referencing Monthly Comparison."""
    ws = wb.create_sheet("CLV Lost to Failures")

    MC = "'Monthly Comparison'"

    # Columns:
    # A: Month
    # B: Failed DTP (A)        C: Failed DTP (C)
    # D: DTP Immed Loss (A)    E: DTP Immed Loss (C)
    # F: DTP Future Loss (A)   G: DTP Future Loss (C)
    # H: Failed Ren (A)        I: Failed Ren (C)
    # J: Ren Loss (A)          K: Ren Loss (C)
    # L: Total CLV Lost (A)    M: Total CLV Lost (C)

    headers = [
        "Month",
        f"Failed DTP ({GW_A})", f"Failed DTP ({GW_C})",
        f"DTP Immediate Loss ({GW_A})", f"DTP Immediate Loss ({GW_C})",
        f"DTP Future Loss ({GW_A})", f"DTP Future Loss ({GW_C})",
        f"Failed Ren ({GW_A})", f"Failed Ren ({GW_C})",
        f"Ren Loss ({GW_A})", f"Ren Loss ({GW_C})",
        f"Total CLV Lost ({GW_A})", f"Total CLV Lost ({GW_C})",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))

    first_data = 2
    last_data = first_data + num_months - 1
    tr = last_data + 1

    for r in range(first_data, last_data + 1):
        mc_r = r  # same row alignment

        # A: Month
        _cell(ws, r, 1, f"={MC}!A{mc_r}")

        # B,C: Failed DTP (from MC cols H,I)
        _cell(ws, r, 2, f"={MC}!H{mc_r}", INT)
        _cell(ws, r, 3, f"={MC}!I{mc_r}", INT)

        # D,E: DTP Immediate Loss = Failed DTP * Avg Txn
        _cell(ws, r, 4, f"=B{r}*{P_AVG_TXN}", MONEY)
        _cell(ws, r, 5, f"=C{r}*{P_AVG_TXN}", MONEY)

        # F,G: DTP Future Loss = Failed DTP * (Lifetime-1) * Avg Txn
        _cell(ws, r, 6, f"=B{r}*({P_LIFETIME}-1)*{P_AVG_TXN}", MONEY)
        _cell(ws, r, 7, f"=C{r}*({P_LIFETIME}-1)*{P_AVG_TXN}", MONEY)

        # H,I: Failed Ren (from MC cols P,Q)
        _cell(ws, r, 8, f"={MC}!P{mc_r}", INT)
        _cell(ws, r, 9, f"={MC}!Q{mc_r}", INT)

        # J,K: Ren Loss = Failed Ren * Avg Txn
        _cell(ws, r, 10, f"=H{r}*{P_AVG_TXN}", MONEY)
        _cell(ws, r, 11, f"=I{r}*{P_AVG_TXN}", MONEY)

        # L,M: Total CLV Lost = DTP Immed + DTP Future + Ren Loss
        _cell(ws, r, 12, f"=D{r}+F{r}+J{r}", MONEY)
        _cell(ws, r, 13, f"=E{r}+G{r}+K{r}", MONEY)

    # Totals row
    _cell(ws, tr, 1, f"{num_months}-Month Total")
    for ci in range(2, 14):
        cl = get_column_letter(ci)
        fmt = MONEY if ci >= 4 else INT
        _cell(ws, tr, ci, f"=SUM({cl}{first_data}:{cl}{last_data})", fmt)

    _style_totals_row(ws, tr, len(headers))
    _set_col_widths(ws, [14] + [22] * 12)
    return ws


def _build_what_if(wb, num_months):
    """What-If: all combined volume through one gateway, using formulas."""
    ws = wb.create_sheet("What-If Analysis")

    RD = "'Raw Data'"

    # Columns:
    # A: Month
    # B: Combined DTP Vol      C: Combined Ren Vol
    # -- All via Adyen --
    # D: DTP SR (A)            E: Ren SR (A)
    # F: Succ DTP (A)          G: Succ Ren (A)
    # H: Fail DTP (A)          I: Fail Ren (A)
    # J: Total Succ (A)
    # K: Gross (A)             L: GW Cost (A)
    # M: Refunds (A)           N: CB (A)
    # O: Net (A)               P: CLV Lost (A)
    # -- All via Cybersource --
    # Q: DTP SR (C)            R: Ren SR (C)
    # S: Succ DTP (C)          T: Succ Ren (C)
    # U: Fail DTP (C)          V: Fail Ren (C)
    # W: Total Succ (C)
    # X: Gross (C)             Y: GW Cost (C)
    # Z: Refunds (C)           AA: CB (C)
    # AB: Net (C)              AC: CLV Lost (C)
    # -- Deltas --
    # AD: Net Delta (A-C)      AE: CLV Lost Delta (A-C)

    headers = [
        "Month", "Combined DTP Vol", "Combined Ren Vol",
        f"DTP SR ({GW_A})", f"Ren SR ({GW_A})",
        f"Succ DTP ({GW_A})", f"Succ Ren ({GW_A})",
        f"Fail DTP ({GW_A})", f"Fail Ren ({GW_A})",
        f"Total Succ ({GW_A})",
        f"Gross ({GW_A})", f"GW Cost ({GW_A})",
        f"Refunds ({GW_A})", f"CB ({GW_A})",
        f"Net ({GW_A})", f"CLV Lost ({GW_A})",
        f"DTP SR ({GW_C})", f"Ren SR ({GW_C})",
        f"Succ DTP ({GW_C})", f"Succ Ren ({GW_C})",
        f"Fail DTP ({GW_C})", f"Fail Ren ({GW_C})",
        f"Total Succ ({GW_C})",
        f"Gross ({GW_C})", f"GW Cost ({GW_C})",
        f"Refunds ({GW_C})", f"CB ({GW_C})",
        f"Net ({GW_C})", f"CLV Lost ({GW_C})",
        "Net Delta (A-C)", "CLV Lost Delta (A-C)",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))

    first_data = 2
    last_data = first_data + num_months - 1
    tr = last_data + 1

    for r in range(first_data, last_data + 1):
        rd_r = r
        _cell(ws, r, 1, f"={RD}!A{rd_r}")

        # B: Combined DTP = Vol_DTP_A + Vol_DTP_C
        _cell(ws, r, 2, f"={RD}!F{rd_r}+{RD}!G{rd_r}", INT)
        # C: Combined Ren
        _cell(ws, r, 3, f"={RD}!H{rd_r}+{RD}!I{rd_r}", INT)

        # --- Adyen side (all combined volume at Adyen rates) ---
        _cell(ws, r, 4, f"={RD}!B{rd_r}", PCT)     # DTP SR
        _cell(ws, r, 5, f"={RD}!D{rd_r}", PCT)     # Ren SR
        _cell(ws, r, 6, f"=ROUND(B{r}*D{r},0)", INT)  # Succ DTP
        _cell(ws, r, 7, f"=ROUND(C{r}*E{r},0)", INT)  # Succ Ren
        _cell(ws, r, 8, f"=B{r}-F{r}", INT)            # Fail DTP
        _cell(ws, r, 9, f"=C{r}-G{r}", INT)            # Fail Ren
        _cell(ws, r, 10, f"=F{r}+G{r}", INT)           # Total Succ
        _cell(ws, r, 11, f"=J{r}*{P_AVG_TXN}", MONEY)                      # Gross
        _cell(ws, r, 12, f"=K{r}*{P_COST_A}", MONEY)                       # GW Cost
        _cell(ws, r, 13, f"=F{r}*{RD}!J{rd_r}+G{r}*{RD}!L{rd_r}", MONEY)  # Refunds
        _cell(ws, r, 14, f"=F{r}*{RD}!N{rd_r}+G{r}*{RD}!P{rd_r}", MONEY)  # CB
        _cell(ws, r, 15, f"=K{r}-L{r}-M{r}-N{r}", MONEY)                   # Net
        _cell(ws, r, 16, f"=H{r}*{P_LIFETIME}*{P_AVG_TXN}+I{r}*{P_AVG_TXN}", MONEY)  # CLV Lost

        # --- Cybersource side ---
        _cell(ws, r, 17, f"={RD}!C{rd_r}", PCT)
        _cell(ws, r, 18, f"={RD}!E{rd_r}", PCT)
        _cell(ws, r, 19, f"=ROUND(B{r}*Q{r},0)", INT)
        _cell(ws, r, 20, f"=ROUND(C{r}*R{r},0)", INT)
        _cell(ws, r, 21, f"=B{r}-S{r}", INT)
        _cell(ws, r, 22, f"=C{r}-T{r}", INT)
        _cell(ws, r, 23, f"=S{r}+T{r}", INT)
        _cell(ws, r, 24, f"=W{r}*{P_AVG_TXN}", MONEY)
        _cell(ws, r, 25, f"=X{r}*{P_COST_C}", MONEY)
        _cell(ws, r, 26, f"=S{r}*{RD}!K{rd_r}+T{r}*{RD}!M{rd_r}", MONEY)
        _cell(ws, r, 27, f"=S{r}*{RD}!O{rd_r}+T{r}*{RD}!Q{rd_r}", MONEY)
        _cell(ws, r, 28, f"=X{r}-Y{r}-Z{r}-AA{r}", MONEY)
        _cell(ws, r, 29, f"=U{r}*{P_LIFETIME}*{P_AVG_TXN}+V{r}*{P_AVG_TXN}", MONEY)

        # Deltas
        _cell(ws, r, 30, f"=O{r}-AB{r}", MONEY)
        _cell(ws, r, 31, f"=P{r}-AC{r}", MONEY)

    # Totals row
    _cell(ws, tr, 1, f"{num_months}-Month Total")
    sum_cols_int = [2, 3, 6, 7, 8, 9, 10, 19, 20, 21, 22, 23]
    sum_cols_money = [11, 12, 13, 14, 15, 16, 24, 25, 26, 27, 28, 29, 30, 31]
    for ci in sum_cols_int:
        cl = get_column_letter(ci)
        _cell(ws, tr, ci, f"=SUM({cl}{first_data}:{cl}{last_data})", INT)
    for ci in sum_cols_money:
        cl = get_column_letter(ci)
        _cell(ws, tr, ci, f"=SUM({cl}{first_data}:{cl}{last_data})", MONEY)
    # SR totals are weighted averages
    _cell(ws, tr, 4, f"=IF(B{tr}=0,0,F{tr}/B{tr})", PCT)
    _cell(ws, tr, 5, f"=IF(C{tr}=0,0,G{tr}/C{tr})", PCT)
    _cell(ws, tr, 17, f"=IF(B{tr}=0,0,S{tr}/B{tr})", PCT)
    _cell(ws, tr, 18, f"=IF(C{tr}=0,0,T{tr}/C{tr})", PCT)

    _style_totals_row(ws, tr, len(headers))
    _set_col_widths(ws, [14] + [18] * (len(headers) - 1))
    return ws


def _build_volume_shift(wb, num_months, scenarios):
    """Volume-Shift scenarios — one block per scenario, all formulas."""
    ws = wb.create_sheet("Volume-Shift Scenarios")

    RD = "'Raw Data'"

    # Each scenario block has:
    #  Header row (scenario label + Adyen %)
    #  Column headers row
    #  num_months data rows
    #  Totals row
    #  Blank row

    block_headers = [
        "Month",
        f"{GW_A} %", f"{GW_C} %",
        "Total DTP", "Total Ren",
        f"DTP Vol ({GW_A})", f"DTP Vol ({GW_C})",
        f"Ren Vol ({GW_A})", f"Ren Vol ({GW_C})",
        f"Succ DTP ({GW_A})", f"Succ DTP ({GW_C})",
        f"Succ Ren ({GW_A})", f"Succ Ren ({GW_C})",
        f"Fail DTP ({GW_A})", f"Fail DTP ({GW_C})",
        f"Fail Ren ({GW_A})", f"Fail Ren ({GW_C})",
        f"Gross ({GW_A})", f"Gross ({GW_C})",
        f"GW Cost ({GW_A})", f"GW Cost ({GW_C})",
        f"Refunds ({GW_A})", f"Refunds ({GW_C})",
        f"CB ({GW_A})", f"CB ({GW_C})",
        f"Net ({GW_A})", f"Net ({GW_C})",
        f"CLV Lost ({GW_A})", f"CLV Lost ({GW_C})",
        "Combined Net", "Combined CLV Lost",
    ]
    num_cols = len(block_headers)

    current_row = 1

    # Summary table location (we'll fill after blocks)
    summary_start = None
    scenario_totals_rows = []

    for si, pct_a in enumerate(scenarios):
        pct_c = 1.0 - pct_a
        label = f"{pct_a:.0%} {GW_A} / {pct_c:.0%} {GW_C}"

        # Scenario title row
        title_cell = ws.cell(row=current_row, column=1, value=label)
        title_cell.font = Font(bold=True, size=12)
        current_row += 1

        # Column headers
        for ci, h in enumerate(block_headers, 1):
            ws.cell(row=current_row, column=ci, value=h)
        _style_header(ws, current_row, num_cols)
        hdr_row = current_row
        current_row += 1

        first_data = current_row
        last_data = first_data + num_months - 1

        for r in range(first_data, last_data + 1):
            rd_r = r - first_data + 2  # Raw Data row (starts at 2)

            _cell(ws, r, 1, f"={RD}!A{rd_r}")

            # B,C: Adyen/Cyber percentage (static values for this scenario)
            _cell(ws, r, 2, pct_a, PCT)
            _cell(ws, r, 3, pct_c, PCT)

            # D: Total DTP = Raw DTP_A + Raw DTP_C
            _cell(ws, r, 4, f"={RD}!F{rd_r}+{RD}!G{rd_r}", INT)
            # E: Total Ren
            _cell(ws, r, 5, f"={RD}!H{rd_r}+{RD}!I{rd_r}", INT)

            # F,G: Redistributed DTP vol
            _cell(ws, r, 6, f"=ROUND(D{r}*B{r},0)", INT)
            _cell(ws, r, 7, f"=D{r}-F{r}", INT)

            # H,I: Redistributed Ren vol
            _cell(ws, r, 8, f"=ROUND(E{r}*B{r},0)", INT)
            _cell(ws, r, 9, f"=E{r}-H{r}", INT)

            # J,K: Succ DTP = ROUND(DTP_vol * SR)
            _cell(ws, r, 10, f"=ROUND(F{r}*{RD}!B{rd_r},0)", INT)
            _cell(ws, r, 11, f"=ROUND(G{r}*{RD}!C{rd_r},0)", INT)

            # L,M: Succ Ren
            _cell(ws, r, 12, f"=ROUND(H{r}*{RD}!D{rd_r},0)", INT)
            _cell(ws, r, 13, f"=ROUND(I{r}*{RD}!E{rd_r},0)", INT)

            # N,O: Fail DTP
            _cell(ws, r, 14, f"=F{r}-J{r}", INT)
            _cell(ws, r, 15, f"=G{r}-K{r}", INT)

            # P,Q: Fail Ren
            _cell(ws, r, 16, f"=H{r}-L{r}", INT)
            _cell(ws, r, 17, f"=I{r}-M{r}", INT)

            # R,S: Gross = (Succ_DTP + Succ_Ren) * Avg Txn
            _cell(ws, r, 18, f"=(J{r}+L{r})*{P_AVG_TXN}", MONEY)
            _cell(ws, r, 19, f"=(K{r}+M{r})*{P_AVG_TXN}", MONEY)

            # T,U: GW Cost
            _cell(ws, r, 20, f"=R{r}*{P_COST_A}", MONEY)
            _cell(ws, r, 21, f"=S{r}*{P_COST_C}", MONEY)

            # V,W: Refunds
            _cell(ws, r, 22, f"=J{r}*{RD}!J{rd_r}+L{r}*{RD}!L{rd_r}", MONEY)
            _cell(ws, r, 23, f"=K{r}*{RD}!K{rd_r}+M{r}*{RD}!M{rd_r}", MONEY)

            # X,Y: Chargebacks
            _cell(ws, r, 24, f"=J{r}*{RD}!N{rd_r}+L{r}*{RD}!P{rd_r}", MONEY)
            _cell(ws, r, 25, f"=K{r}*{RD}!O{rd_r}+M{r}*{RD}!Q{rd_r}", MONEY)

            # Z,AA: Net = Gross - Cost - Refunds - CB
            _cell(ws, r, 26, f"=R{r}-T{r}-V{r}-X{r}", MONEY)
            _cell(ws, r, 27, f"=S{r}-U{r}-W{r}-Y{r}", MONEY)

            # AB,AC: CLV Lost = Fail_DTP * Lifetime * AvgTxn + Fail_Ren * AvgTxn
            _cell(ws, r, 28, f"=N{r}*{P_LIFETIME}*{P_AVG_TXN}+P{r}*{P_AVG_TXN}", MONEY)
            _cell(ws, r, 29, f"=O{r}*{P_LIFETIME}*{P_AVG_TXN}+Q{r}*{P_AVG_TXN}", MONEY)

            # AD: Combined Net
            _cell(ws, r, 30, f"=Z{r}+AA{r}", MONEY)
            # AE: Combined CLV Lost
            _cell(ws, r, 31, f"=AB{r}+AC{r}", MONEY)

        # Totals row
        tr = last_data + 1
        _cell(ws, tr, 1, f"{num_months}-Month Total")
        sum_money = [18,19,20,21,22,23,24,25,26,27,28,29,30,31]
        sum_int = [4,5,6,7,8,9,10,11,12,13,14,15,16,17]
        for ci in sum_int:
            cl = get_column_letter(ci)
            _cell(ws, tr, ci, f"=SUM({cl}{first_data}:{cl}{last_data})", INT)
        for ci in sum_money:
            cl = get_column_letter(ci)
            _cell(ws, tr, ci, f"=SUM({cl}{first_data}:{cl}{last_data})", MONEY)
        _cell(ws, tr, 2, pct_a, PCT)
        _cell(ws, tr, 3, pct_c, PCT)

        _style_totals_row(ws, tr, num_cols)
        scenario_totals_rows.append(tr)

        current_row = tr + 2  # blank row between blocks

    # --- Summary table at the bottom ---
    summary_start = current_row + 1
    ws.cell(row=summary_start, column=1, value="VOLUME-SHIFT SUMMARY").font = Font(bold=True, size=12)
    summary_start += 1

    sum_headers = ["Scenario", "Combined Net Revenue", "Combined CLV Lost",
                   "Net - CLV Lost", "Delta Net vs Baseline", "Delta CLV Lost vs Baseline"]
    for ci, h in enumerate(sum_headers, 1):
        ws.cell(row=summary_start, column=ci, value=h)
    _style_header(ws, summary_start, len(sum_headers))

    baseline_tr = scenario_totals_rows[0]
    ad_col = get_column_letter(30)  # AD = Combined Net
    ae_col = get_column_letter(31)  # AE = Combined CLV Lost

    for si, (pct_a, tr) in enumerate(zip(scenarios, scenario_totals_rows)):
        pct_c = 1.0 - pct_a
        sr = summary_start + 1 + si
        label = f"{pct_a:.0%} {GW_A} / {pct_c:.0%} {GW_C}"
        _cell(ws, sr, 1, label)
        _cell(ws, sr, 2, f"={ad_col}{tr}", MONEY)  # Combined Net
        _cell(ws, sr, 3, f"={ae_col}{tr}", MONEY)   # Combined CLV Lost
        _cell(ws, sr, 4, f"=B{sr}-C{sr}", MONEY)     # Net - CLV Lost
        if si == 0:
            _cell(ws, sr, 5, 0, MONEY)
            _cell(ws, sr, 6, 0, MONEY)
        else:
            baseline_sr = summary_start + 1
            _cell(ws, sr, 5, f"=B{sr}-B{baseline_sr}", MONEY)
            _cell(ws, sr, 6, f"=C{sr}-C{baseline_sr}", MONEY)

    _set_col_widths(ws, [14] + [18] * (num_cols - 1))
    return ws


# ═══════════════════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════════════════

def main():
    input_path = Path(INPUT_DIR) / INPUT_FILENAME
    if not input_path.exists():
        print(f"Error: {input_path} not found")
        sys.exit(1)

    out_dir = Path(OUTPUT_DIR)
    os.makedirs(out_dir, exist_ok=True)

    print(f"Loading data from {input_path} ...")
    params, month_list = load_data(str(input_path))
    num_months = len(month_list)
    print(f"  {GW_A} cost: {params['cost_a']:.1%}")
    print(f"  {GW_C} cost: {params['cost_c']:.1%}")
    print(f"  Avg transaction: ${params['avg_txn']:.2f}")
    print(f"  Lifetime: {params['lifetime']} months")
    print(f"  Months loaded: {num_months} ({month_list[0]['label']} -- {month_list[-1]['label']})")

    scenarios = [0.10, 0.20, 0.25]

    wb = Workbook()
    wb.remove(wb.active)

    print("  Building Parameters sheet...")
    _build_parameters(wb, params)

    print("  Building Raw Data sheet...")
    _build_raw_data(wb, month_list)

    print("  Building Monthly Comparison sheet (formulas)...")
    _, mc_first, mc_last, mc_totals = _build_monthly_comparison(wb, num_months)

    print("  Building CLV Lost to Failures sheet (formulas)...")
    _build_clv_lost(wb, num_months, mc_first, mc_last)

    print("  Building What-If Analysis sheet (formulas)...")
    _build_what_if(wb, num_months)

    print("  Building Volume-Shift Scenarios sheet (formulas)...")
    _build_volume_shift(wb, num_months, scenarios)

    output_path = out_dir / "gateway_comparison_formulas.xlsx"
    wb.save(str(output_path))
    print(f"\n[OK] Formula-based Excel workbook saved to {output_path}")
    print("     All derived cells contain Excel formulas — edit Parameters to recalculate.")


if __name__ == "__main__":
    main()
