#!/usr/bin/env python3
"""Gateway Performance Comparison Tool — Adyen vs Cybersource

Reads gateway performance data from an xlsx file and produces a comprehensive
comparison between Adyen and Cybersource, considering success rates, costs,
refunds, chargebacks, customer lifetime value, and volume-shift scenarios.

Usage:
    python3 scripts/gateway_compare.py /path/to/gateway_performance_test.xlsx

Output:
    - Console report with monthly and period-total comparisons
    - Excel workbook (gateway_comparison_results.xlsx) with detailed sheets
    - 6 presentation-ready PNG charts in a charts/ directory

Excel input format:
    B1: Adyen cost rate           B2: Cybersource cost rate
    B3: Avg transaction $         B4: Customer lifetime (months)
    Row 6: Column headers         Rows 7+: Monthly data
    Columns B-E: Success rates (DTP-Adyen, DTP-Cyber, Ren-Adyen, Ren-Cyber)
    Columns F-I: Volumes          Columns J-M: Refunds per txn
    Columns N-Q: Chargebacks per txn
"""

import sys
import os
import argparse
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.ticker as mticker
    import numpy as np
    HAS_MPL = True
except ImportError:
    HAS_MPL = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

GW_A = "Adyen"
GW_C = "Cybersource"

CLR_ADYEN = "#1f77b4"
CLR_CYBER = "#ff7f0e"
CLR_ADYEN_LT = "#aec7e8"
CLR_CYBER_LT = "#ffbb78"
CLR_GREEN = "#2ca02c"
CLR_RED = "#d62728"
CLR_GRAY = "#999999"
CLR_PURPLE = "#9467bd"
FIG_W, FIG_H = 14, 7
DPI = 150

NS_STRICT = "http://purl.oclc.org/ooxml/spreadsheetml/main"
NS_TRANS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

VOLUME_SHIFT_SCENARIOS = [0.10, 0.20, 0.25]


# ═══════════════════════════════════════════════════════════════════════════
#  DATA LOADING  (handles strict OOXML that openpyxl can't always read)
# ═══════════════════════════════════════════════════════════════════════════

def _detect_ns(xml_bytes):
    snippet = xml_bytes[:2000].decode("utf-8", errors="ignore")
    return NS_STRICT if NS_STRICT in snippet else NS_TRANS


def _parse_shared_strings(xml_bytes, ns):
    root = ET.fromstring(xml_bytes)
    return [
        (si.find(f"{{{ns}}}t").text or "")
        if si.find(f"{{{ns}}}t") is not None
        else ""
        for si in root.findall(f"{{{ns}}}si")
    ]


def _col_to_idx(col):
    idx = 0
    for ch in col:
        idx = idx * 26 + (ord(ch.upper()) - 64)
    return idx - 1


def _parse_sheet(xml_bytes, ns, shared):
    root = ET.fromstring(xml_bytes)
    rows = {}
    for row_el in root.iter(f"{{{ns}}}row"):
        rn = int(row_el.get("r"))
        cells = {}
        for c in row_el.findall(f"{{{ns}}}c"):
            ref = c.get("r")
            col = _col_to_idx("".join(ch for ch in ref if ch.isalpha()))
            v = c.find(f"{{{ns}}}v")
            if v is None or v.text is None:
                continue
            raw = v.text
            ct = c.get("t", "")
            if ct == "s":
                cells[col] = shared[int(raw)]
            elif ct == "d":
                cells[col] = raw
            else:
                try:
                    cells[col] = int(raw) if "." not in raw else float(raw)
                except ValueError:
                    cells[col] = raw
        rows[rn] = cells
    return rows


def load_data(filepath):
    """Return (params, month_list) parsed from the xlsx."""
    with zipfile.ZipFile(filepath) as zf:
        ss_xml = zf.read("xl/sharedStrings.xml")
        ns = _detect_ns(ss_xml)
        shared = _parse_shared_strings(ss_xml, ns)
        sh_xml = zf.read("xl/worksheets/sheet1.xml")
        rows = _parse_sheet(sh_xml, _detect_ns(sh_xml), shared)

    params = {
        "cost_a": float(rows.get(1, {}).get(1, 0)),
        "cost_c": float(rows.get(2, {}).get(1, 0)),
        "avg_txn": float(rows.get(3, {}).get(1, 0)),
        "lifetime": int(rows.get(4, {}).get(1, 0)),
    }

    months = []
    for rn in sorted(rows):
        if rn < 7:
            continue
        r = rows[rn]
        if 0 not in r:
            continue
        ds = r[0]

        if isinstance(ds, str):
            if "-" in ds and len(ds) == 10:
                try:
                    dt = datetime.strptime(ds, "%Y-%m-%d")
                    label = dt.strftime("%b %Y")
                except ValueError:
                    label = ds
            else:
                label = str(ds)
        elif isinstance(ds, (int, float)):
            label = str(ds)
        else:
            label = str(ds)

        months.append({
            "label": label,
            "sr_ci_a": float(r.get(1, 0)), "sr_ci_c": float(r.get(2, 0)),
            "sr_ren_a": float(r.get(3, 0)), "sr_ren_c": float(r.get(4, 0)),
            "vol_ci_a": int(r.get(5, 0)), "vol_ci_c": int(r.get(6, 0)),
            "vol_ren_a": int(r.get(7, 0)), "vol_ren_c": int(r.get(8, 0)),
            "ref_ci_a": float(r.get(9, 0)), "ref_ci_c": float(r.get(10, 0)),
            "ref_ren_a": float(r.get(11, 0)), "ref_ren_c": float(r.get(12, 0)),
            "cb_ci_a": float(r.get(13, 0)), "cb_ci_c": float(r.get(14, 0)),
            "cb_ren_a": float(r.get(15, 0)), "cb_ren_c": float(r.get(16, 0)),
        })
    return params, months


# ═══════════════════════════════════════════════════════════════════════════
#  CALCULATIONS
# ═══════════════════════════════════════════════════════════════════════════

def _calc_gateway(md, gw, params):
    """Compute all metrics for one gateway in one month."""
    avg = params["avg_txn"]
    lt = params["lifetime"]
    cr = params[f"cost_{gw}"]
    sr_ci = md[f"sr_ci_{gw}"]
    sr_ren = md[f"sr_ren_{gw}"]
    vol_ci = md[f"vol_ci_{gw}"]
    vol_ren = md[f"vol_ren_{gw}"]
    ref_ci = md[f"ref_ci_{gw}"]
    ref_ren = md[f"ref_ren_{gw}"]
    cb_ci = md[f"cb_ci_{gw}"]
    cb_ren = md[f"cb_ren_{gw}"]

    succ_ci = round(vol_ci * sr_ci)
    fail_ci = vol_ci - succ_ci
    succ_ren = round(vol_ren * sr_ren)
    fail_ren = vol_ren - succ_ren
    total_vol = vol_ci + vol_ren
    total_succ = succ_ci + succ_ren
    overall_sr = total_succ / total_vol if total_vol else 0

    gross = total_succ * avg
    gw_cost = gross * cr
    refunds = succ_ci * ref_ci + succ_ren * ref_ren
    cbacks = succ_ci * cb_ci + succ_ren * cb_ren
    net = gross - gw_cost - refunds - cbacks

    lost_imm_ci = fail_ci * avg
    lost_fut_ci = fail_ci * (lt - 1) * avg
    lost_ren_rev = fail_ren * avg
    clv_impact = lost_imm_ci + lost_fut_ci + lost_ren_rev

    exp_ci = sr_ci * (avg - avg * cr - ref_ci - cb_ci)
    exp_ren = sr_ren * (avg - avg * cr - ref_ren - cb_ren)

    return {
        "sr_ci": sr_ci, "sr_ren": sr_ren,
        "vol_ci": vol_ci, "vol_ren": vol_ren,
        "succ_ci": succ_ci, "fail_ci": fail_ci,
        "succ_ren": succ_ren, "fail_ren": fail_ren,
        "total_vol": total_vol, "total_succ": total_succ, "overall_sr": overall_sr,
        "gross": gross, "gw_cost": gw_cost,
        "refunds": refunds, "chargebacks": cbacks, "net": net,
        "lost_imm_ci": lost_imm_ci, "lost_fut_ci": lost_fut_ci,
        "lost_ren": lost_ren_rev, "clv_impact": clv_impact,
        "net_per_attempt": net / total_vol if total_vol else 0,
        "net_per_succ": net / total_succ if total_succ else 0,
        "exp_ci": exp_ci, "exp_ren": exp_ren,
        "ref_ci": ref_ci, "ref_ren": ref_ren,
        "cb_ci": cb_ci, "cb_ren": cb_ren,
    }


def calculate_all(month_list, params):
    """Return list of monthly result dicts, each containing 'a' and 'c' sub-dicts."""
    results = []
    for md in month_list:
        results.append({
            "label": md["label"],
            "a": _calc_gateway(md, "a", params),
            "c": _calc_gateway(md, "c", params),
        })
    return results


def aggregate_totals(monthly):
    """Sum monthly results into period totals and recompute rates."""
    SUM_KEYS = [
        "vol_ci", "vol_ren", "succ_ci", "fail_ci", "succ_ren", "fail_ren",
        "total_vol", "total_succ", "gross", "gw_cost", "refunds", "chargebacks",
        "net", "lost_imm_ci", "lost_fut_ci", "lost_ren", "clv_impact",
    ]
    totals = {"label": f"{len(monthly)}-Month Total"}
    for gw in ("a", "c"):
        t = {}
        for k in SUM_KEYS:
            t[k] = sum(m[gw][k] for m in monthly)
        t["sr_ci"] = t["succ_ci"] / t["vol_ci"] if t["vol_ci"] else 0
        t["sr_ren"] = t["succ_ren"] / t["vol_ren"] if t["vol_ren"] else 0
        t["overall_sr"] = t["total_succ"] / t["total_vol"] if t["total_vol"] else 0
        t["net_per_attempt"] = t["net"] / t["total_vol"] if t["total_vol"] else 0
        t["net_per_succ"] = t["net"] / t["total_succ"] if t["total_succ"] else 0
        t["gw_cost_per_succ"] = t["gw_cost"] / t["total_succ"] if t["total_succ"] else 0
        totals[gw] = t
    return totals


def what_if(month_list, params):
    """What-if: all combined volume through each gateway, per month + total."""
    avg = params["avg_txn"]
    lt = params["lifetime"]
    wi_months = []
    for md in month_list:
        tot_ci = md["vol_ci_a"] + md["vol_ci_c"]
        tot_ren = md["vol_ren_a"] + md["vol_ren_c"]
        wi = {"label": md["label"], "tot_ci": tot_ci, "tot_ren": tot_ren}
        for gw in ("a", "c"):
            cr = params[f"cost_{gw}"]
            sr_ci = md[f"sr_ci_{gw}"]
            sr_ren = md[f"sr_ren_{gw}"]
            ref_ci = md[f"ref_ci_{gw}"]
            ref_ren = md[f"ref_ren_{gw}"]
            cb_ci = md[f"cb_ci_{gw}"]
            cb_ren = md[f"cb_ren_{gw}"]
            s_ci = round(tot_ci * sr_ci)
            s_ren = round(tot_ren * sr_ren)
            f_ci = tot_ci - s_ci
            f_ren = tot_ren - s_ren
            ts = s_ci + s_ren
            g = ts * avg
            gc = g * cr
            ref = s_ci * ref_ci + s_ren * ref_ren
            cb = s_ci * cb_ci + s_ren * cb_ren
            n = g - gc - ref - cb
            clv = f_ci * lt * avg + f_ren * avg
            wi[gw] = {
                "sr_ci": sr_ci, "sr_ren": sr_ren,
                "succ_ci": s_ci, "succ_ren": s_ren,
                "fail_ci": f_ci, "fail_ren": f_ren,
                "total_succ": ts,
                "gross": g, "gw_cost": gc,
                "refunds": ref, "chargebacks": cb,
                "net": n, "clv": clv,
            }
        wi_months.append(wi)

    wi_tot = {"label": "Total", "tot_ci": 0, "tot_ren": 0}
    for gw in ("a", "c"):
        t = {}
        for k in ("succ_ci", "succ_ren", "fail_ci", "fail_ren", "total_succ",
                   "gross", "gw_cost", "refunds", "chargebacks", "net", "clv"):
            t[k] = sum(w[gw][k] for w in wi_months)
        wi_tot[gw] = t
    wi_tot["tot_ci"] = sum(w["tot_ci"] for w in wi_months)
    wi_tot["tot_ren"] = sum(w["tot_ren"] for w in wi_months)
    return wi_months, wi_tot


def volume_shift(month_list, params, scenarios=None):
    """Model DTP volume redistribution scenarios between gateways.

    For each scenario percentage (Adyen's share of DTP), redistributes DTP and
    renewal volumes, applies each gateway's observed rates, and computes financials.
    Renewals are scaled to match the DTP split ratio since customers stay on
    their assigned gateway for life.
    """
    if scenarios is None:
        scenarios = VOLUME_SHIFT_SCENARIOS
    avg = params["avg_txn"]
    lt = params["lifetime"]

    results = []
    for pct_a in scenarios:
        pct_c = 1.0 - pct_a
        scenario_months = []
        for md in month_list:
            tot_ci = md["vol_ci_a"] + md["vol_ci_c"]
            tot_ren = md["vol_ren_a"] + md["vol_ren_c"]

            new_ci_a = round(tot_ci * pct_a)
            new_ci_c = tot_ci - new_ci_a
            new_ren_a = round(tot_ren * pct_a)
            new_ren_c = tot_ren - new_ren_a

            sm = {"label": md["label"], "tot_ci": tot_ci, "tot_ren": tot_ren}
            for gw, ci_vol, ren_vol in [("a", new_ci_a, new_ren_a),
                                         ("c", new_ci_c, new_ren_c)]:
                cr = params[f"cost_{gw}"]
                sr_ci = md[f"sr_ci_{gw}"]
                sr_ren = md[f"sr_ren_{gw}"]
                ref_ci = md[f"ref_ci_{gw}"]
                ref_ren = md[f"ref_ren_{gw}"]
                cb_ci = md[f"cb_ci_{gw}"]
                cb_ren = md[f"cb_ren_{gw}"]

                s_ci = round(ci_vol * sr_ci)
                f_ci = ci_vol - s_ci
                s_ren = round(ren_vol * sr_ren)
                f_ren = ren_vol - s_ren
                ts = s_ci + s_ren
                g = ts * avg
                gc = g * cr
                ref = s_ci * ref_ci + s_ren * ref_ren
                cb = s_ci * cb_ci + s_ren * cb_ren
                n = g - gc - ref - cb
                clv = f_ci * lt * avg + f_ren * avg

                sm[gw] = {
                    "vol_ci": ci_vol, "vol_ren": ren_vol,
                    "sr_ci": sr_ci, "sr_ren": sr_ren,
                    "succ_ci": s_ci, "fail_ci": f_ci,
                    "succ_ren": s_ren, "fail_ren": f_ren,
                    "total_succ": ts,
                    "gross": g, "gw_cost": gc,
                    "refunds": ref, "chargebacks": cb,
                    "net": n, "clv": clv,
                }
            sm["combined_net"] = sm["a"]["net"] + sm["c"]["net"]
            sm["combined_clv"] = sm["a"]["clv"] + sm["c"]["clv"]
            sm["combined_gross"] = sm["a"]["gross"] + sm["c"]["gross"]
            sm["combined_gw_cost"] = sm["a"]["gw_cost"] + sm["c"]["gw_cost"]
            sm["combined_refunds"] = sm["a"]["refunds"] + sm["c"]["refunds"]
            sm["combined_chargebacks"] = sm["a"]["chargebacks"] + sm["c"]["chargebacks"]
            sm["combined_succ"] = sm["a"]["total_succ"] + sm["c"]["total_succ"]
            sm["combined_fail_ci"] = sm["a"]["fail_ci"] + sm["c"]["fail_ci"]
            scenario_months.append(sm)

        totals = {
            "pct_a": pct_a,
            "pct_c": pct_c,
            "label": f"{pct_a:.0%} {GW_A} / {pct_c:.0%} {GW_C}",
            "months": scenario_months,
            "combined_net": sum(m["combined_net"] for m in scenario_months),
            "combined_clv": sum(m["combined_clv"] for m in scenario_months),
            "combined_gross": sum(m["combined_gross"] for m in scenario_months),
            "combined_gw_cost": sum(m["combined_gw_cost"] for m in scenario_months),
            "combined_refunds": sum(m["combined_refunds"] for m in scenario_months),
            "combined_chargebacks": sum(m["combined_chargebacks"] for m in scenario_months),
            "combined_succ": sum(m["combined_succ"] for m in scenario_months),
            "combined_fail_ci": sum(m["combined_fail_ci"] for m in scenario_months),
            "adyen_net": sum(m["a"]["net"] for m in scenario_months),
            "cyber_net": sum(m["c"]["net"] for m in scenario_months),
            "adyen_clv": sum(m["a"]["clv"] for m in scenario_months),
            "cyber_clv": sum(m["c"]["clv"] for m in scenario_months),
        }
        results.append(totals)
    return results


# ═══════════════════════════════════════════════════════════════════════════
#  CONSOLE REPORT
# ═══════════════════════════════════════════════════════════════════════════

def _f(n):
    return f"{n:,.0f}"

def _fd(n):
    return f"${n:,.2f}"

def _fp(n):
    return f"{n:.2%}"


def _print_month_block(r):
    a, c = r["a"], r["c"]
    W = 120
    hdr = f"{'Metric':<48} {GW_A:>20} {GW_C:>20} {'Diff (A-C)':>20}"

    print(f"\n{'─' * W}")
    print(f"  {r['label']}")
    print(f"{'─' * W}")
    print(hdr)
    print("-" * W)

    def row(label, va, vc, fmt=_f):
        print(f"{label:<48} {fmt(va):>20} {fmt(vc):>20} {fmt(va - vc):>20}")

    row("DTP Attempts", a["vol_ci"], c["vol_ci"])
    row("DTP Success Rate", a["sr_ci"], c["sr_ci"], _fp)
    row("DTP Successful Txns", a["succ_ci"], c["succ_ci"])
    row("DTP Failed Txns", a["fail_ci"], c["fail_ci"])
    print()
    row("Renewal Attempts", a["vol_ren"], c["vol_ren"])
    row("Renewal Success Rate", a["sr_ren"], c["sr_ren"], _fp)
    row("Renewal Successful Txns", a["succ_ren"], c["succ_ren"])
    row("Renewal Failed Txns", a["fail_ren"], c["fail_ren"])
    print()
    row("Total Attempts", a["total_vol"], c["total_vol"])
    row("Total Successful Txns", a["total_succ"], c["total_succ"])
    row("Overall Success Rate", a["overall_sr"], c["overall_sr"], _fp)
    print()
    row("Gross Revenue", a["gross"], c["gross"], _fd)
    row("Gateway Processing Cost", a["gw_cost"], c["gw_cost"], _fd)
    row("Refunds", a["refunds"], c["refunds"], _fd)
    row("Chargebacks", a["chargebacks"], c["chargebacks"], _fd)
    row("Net Revenue", a["net"], c["net"], _fd)
    print()
    row("Net Revenue per Attempt", a["net_per_attempt"], c["net_per_attempt"], _fd)
    row("Net Revenue per Successful Txn", a["net_per_succ"], c["net_per_succ"], _fd)


def print_report(monthly, totals, wi_months, wi_totals, vs_results, params):
    W = 120
    avg = params["avg_txn"]
    lt = params["lifetime"]

    # ── Methodology ──
    print("=" * W)
    print("METHODOLOGY & ASSUMPTIONS")
    print("=" * W)
    print(f"""
  This report compares {GW_A} and {GW_C} gateway performance across {len(monthly)} months.

  Parameters:
    {GW_A} cost rate:       {params['cost_a']:.1%} of gross revenue
    {GW_C} cost rate:    {params['cost_c']:.1%} of gross revenue
    Avg transaction amount:  {_fd(avg)}
    Customer lifetime:       {lt} months (1 direct-to-paid + {lt - 1} renewals)
    CLV per customer:        {lt} x {_fd(avg)} = {_fd(lt * avg)}

  Formulas:
    Successful Txns       = Volume x Success Rate
    Gross Revenue         = Successful Txns x Avg Transaction
    Gateway Cost          = Gross Revenue x Cost Rate
    Refunds               = Successful DTP x Refund/DTP + Successful Ren x Refund/Ren
    Chargebacks           = Successful DTP x CB/DTP + Successful Ren x CB/Ren
    Net Revenue           = Gross - Gateway Cost - Refunds - Chargebacks
    CLV Lost (DTP fail)   = Failed DTP x Lifetime x Avg Transaction
    CLV Lost (Ren fail)   = Failed Renewals x Avg Transaction
    Expected Net/Attempt  = Success Rate x (Avg Txn - Avg Txn x Cost Rate - Refund - CB)
""")

    # ── Section 1 ──
    print("=" * W)
    print("SECTION 1: MONTHLY DIRECT FINANCIAL COMPARISON")
    print(f"  {GW_A} cost: {params['cost_a']:.1%}   {GW_C} cost: {params['cost_c']:.1%}"
          f"   Avg txn: {_fd(avg)}   Lifetime: {lt} months")
    print("=" * W)
    for r in monthly:
        _print_month_block(r)

    # ── Period totals ──
    a, c = totals["a"], totals["c"]
    print(f"\n{'=' * W}")
    print(f"  {totals['label']}")
    print(f"{'=' * W}")
    hdr = f"{'Metric':<48} {GW_A:>20} {GW_C:>20} {'Diff (A-C)':>20}"
    print(hdr)
    print("-" * W)

    def row(label, va, vc, fmt=_f):
        print(f"{label:<48} {fmt(va):>20} {fmt(vc):>20} {fmt(va - vc):>20}")

    row("DTP Attempts", a["vol_ci"], c["vol_ci"])
    row("DTP Successful", a["succ_ci"], c["succ_ci"])
    row("DTP Wtd Avg Success Rate", a["sr_ci"], c["sr_ci"], _fp)
    print()
    row("Renewal Attempts", a["vol_ren"], c["vol_ren"])
    row("Renewal Successful", a["succ_ren"], c["succ_ren"])
    row("Renewal Wtd Avg Success Rate", a["sr_ren"], c["sr_ren"], _fp)
    print()
    row("Total Attempts", a["total_vol"], c["total_vol"])
    row("Total Successful", a["total_succ"], c["total_succ"])
    row("Overall Success Rate", a["overall_sr"], c["overall_sr"], _fp)
    print()
    row("Gross Revenue", a["gross"], c["gross"], _fd)
    row("Gateway Processing Cost", a["gw_cost"], c["gw_cost"], _fd)
    row("Refunds", a["refunds"], c["refunds"], _fd)
    row("Chargebacks", a["chargebacks"], c["chargebacks"], _fd)
    row("Net Revenue", a["net"], c["net"], _fd)
    print()
    row("Net Revenue per Attempt", a["net_per_attempt"], c["net_per_attempt"], _fd)
    row("Net Revenue per Successful Txn", a["net_per_succ"], c["net_per_succ"], _fd)
    row("Gateway Cost per Successful Txn", a["gw_cost_per_succ"], c["gw_cost_per_succ"], _fd)

    # ── Section 2 ──
    print(f"\n\n{'=' * W}")
    print("SECTION 2: PER-TRANSACTION EXPECTED VALUE (Normalized per attempt)")
    print("=" * W)
    print(f"\nRemoves volume differences to compare gateway efficiency head-to-head.")
    print(f"Expected Net = Success Rate x (Avg Txn - Gateway Cost - Refund - Chargeback)\n")
    for r in monthly:
        a, c = r["a"], r["c"]
        print(f"  {r['label']}:")
        print(f"    DTP -> {GW_A}: {_fd(a['exp_ci'])}  {GW_C}: {_fd(c['exp_ci'])}"
              f"  Delta(A-C): {_fd(a['exp_ci'] - c['exp_ci'])}"
              f"  {'{} wins'.format(GW_A) if a['exp_ci'] > c['exp_ci'] else '{} wins'.format(GW_C)}")
        print(f"    Ren -> {GW_A}: {_fd(a['exp_ren'])}  {GW_C}: {_fd(c['exp_ren'])}"
              f"  Delta(A-C): {_fd(a['exp_ren'] - c['exp_ren'])}"
              f"  {'{} wins'.format(GW_A) if a['exp_ren'] > c['exp_ren'] else '{} wins'.format(GW_C)}")
        print()

    # ── Section 3 ──
    print(f"{'=' * W}")
    print("SECTION 3: CUSTOMER LIFETIME VALUE (CLV) LOST TO FAILURES")
    print("=" * W)
    print(f"""
  Assumptions:
    - Failed DTP = lost customer (never signs up)
    - Customer lifetime: {lt} months (1 DTP + {lt - 1} renewals)
    - Avg transaction: {_fd(avg)}
    - CLV per customer: {lt} x {_fd(avg)} = {_fd(lt * avg)}
    - CLV lost per failed DTP: {_fd(lt * avg)}
    - Renewal failure loses only {_fd(avg)} immediate revenue
""")
    for r in monthly:
        a, c = r["a"], r["c"]
        print(f"  {r['label']}:")
        print(f"    {'Metric':<50} {GW_A:>18} {GW_C:>18}")
        print(f"    {'-' * 90}")
        print(f"    {'Failed DTP Txns':<50} {_f(a['fail_ci']):>18} {_f(c['fail_ci']):>18}")
        print(f"    {'DTP Immediate Lost Revenue':<50} {_fd(a['lost_imm_ci']):>18} {_fd(c['lost_imm_ci']):>18}")
        print(f"    {'DTP Future Revenue Lost':<50} {_fd(a['lost_fut_ci']):>18} {_fd(c['lost_fut_ci']):>18}")
        print(f"    {'Failed Renewal Txns':<50} {_f(a['fail_ren']):>18} {_f(c['fail_ren']):>18}")
        print(f"    {'Renewal Lost Revenue':<50} {_fd(a['lost_ren']):>18} {_fd(c['lost_ren']):>18}")
        print(f"    {'TOTAL Economic Impact':<50} {_fd(a['clv_impact']):>18} {_fd(c['clv_impact']):>18}")
        print()

    # ── Section 4 ──
    print(f"{'=' * W}")
    print(f"SECTION 4: WHAT-IF -- If ALL combined volume were routed through one gateway")
    print("=" * W)
    print(f"\nApplies each gateway's observed success rates to the combined volume.\n")
    for w in wi_months:
        a, c = w["a"], w["c"]
        cost_a_label = f"All via {GW_A} ({params['cost_a']:.1%})"
        cost_c_label = f"All via {GW_C} ({params['cost_c']:.1%})"
        print(f"  {w['label']}  |  Combined DTP: {_f(w['tot_ci'])}  |  Combined Ren: {_f(w['tot_ren'])}")
        print(f"    {'Metric':<48} {cost_a_label:>18} {cost_c_label:>18} {'Delta (A-C)':>18}")
        print(f"    {'-' * 106}")
        print(f"    {'DTP Success Rate':<48} {_fp(a['sr_ci']):>18} {_fp(c['sr_ci']):>18}")
        print(f"    {'Renewal Success Rate':<48} {_fp(a['sr_ren']):>18} {_fp(c['sr_ren']):>18}")
        print(f"    {'Successful Txns':<48} {_f(a['total_succ']):>18} {_f(c['total_succ']):>18} {_f(a['total_succ'] - c['total_succ']):>18}")
        print(f"    {'Gross Revenue':<48} {_fd(a['gross']):>18} {_fd(c['gross']):>18} {_fd(a['gross'] - c['gross']):>18}")
        print(f"    {'Gateway Cost':<48} {_fd(a['gw_cost']):>18} {_fd(c['gw_cost']):>18} {_fd(a['gw_cost'] - c['gw_cost']):>18}")
        print(f"    {'Refunds':<48} {_fd(a['refunds']):>18} {_fd(c['refunds']):>18} {_fd(a['refunds'] - c['refunds']):>18}")
        print(f"    {'Chargebacks':<48} {_fd(a['chargebacks']):>18} {_fd(c['chargebacks']):>18} {_fd(a['chargebacks'] - c['chargebacks']):>18}")
        print(f"    {'Net Revenue':<48} {_fd(a['net']):>18} {_fd(c['net']):>18} {_fd(a['net'] - c['net']):>18}")
        print(f"    {'CLV Lost to Failures':<48} {_fd(a['clv']):>18} {_fd(c['clv']):>18} {_fd(a['clv'] - c['clv']):>18}")
        print()

    a, c = wi_totals["a"], wi_totals["c"]
    print(f"  {'-' * 106}")
    print(f"  {totals['label'].upper()} WHAT-IF TOTALS:")
    print(f"    {'Total Net Revenue':<48} {_fd(a['net']):>18} {_fd(c['net']):>18} {_fd(a['net'] - c['net']):>18}")
    print(f"    {'Total CLV Lost to Failures':<48} {_fd(a['clv']):>18} {_fd(c['clv']):>18} {_fd(a['clv'] - c['clv']):>18}")
    print(f"    {'Net Revenue minus CLV Lost':<48} {_fd(a['net'] - a['clv']):>18} {_fd(c['net'] - c['clv']):>18} {_fd((a['net'] - a['clv']) - (c['net'] - c['clv'])):>18}")

    # ── Section 5: Volume Shift ──
    print(f"\n\n{'=' * W}")
    print(f"SECTION 5: VOLUME-SHIFT SCENARIO ANALYSIS")
    print(f"  What happens if {GW_A}'s share of direct-to-paid volume increases?")
    print("=" * W)
    print(f"""
  Currently {GW_A} receives ~10% of DTP volume, {GW_C} receives ~90%.
  This section models the impact of shifting to 20% or 25% {GW_A}.

  Methodology:
    - Total DTP volume = {GW_A} DTP + {GW_C} DTP (combined each month)
    - Redistribute DTP according to new split percentage
    - Renewal volume scales to same split (customers stay on their gateway)
    - Apply each gateway's OBSERVED monthly success rates to redistributed volume
    - Compute net revenue and CLV lost to failures for each scenario
""")

    baseline = vs_results[0]
    for vs in vs_results:
        pct_a = vs["pct_a"]
        print(f"  {'─' * 106}")
        print(f"  Scenario: {vs['label']}")
        print(f"  {'─' * 106}")

        for sm in vs["months"]:
            a, c = sm["a"], sm["c"]
            print(f"    {sm['label']}:")
            print(f"      {'':36} {GW_A:>16} {GW_C:>16} {'Combined':>16}")
            print(f"      {'DTP Volume':<36} {_f(a['vol_ci']):>16} {_f(c['vol_ci']):>16} {_f(sm['tot_ci']):>16}")
            print(f"      {'Renewal Volume':<36} {_f(a['vol_ren']):>16} {_f(c['vol_ren']):>16} {_f(sm['tot_ren']):>16}")
            print(f"      {'Successful Txns':<36} {_f(a['total_succ']):>16} {_f(c['total_succ']):>16} {_f(sm['combined_succ']):>16}")
            print(f"      {'Net Revenue':<36} {_fd(a['net']):>16} {_fd(c['net']):>16} {_fd(sm['combined_net']):>16}")
            print(f"      {'CLV Lost to Failures':<36} {_fd(a['clv']):>16} {_fd(c['clv']):>16} {_fd(sm['combined_clv']):>16}")
            print()

        delta_net = vs["combined_net"] - baseline["combined_net"]
        delta_clv = vs["combined_clv"] - baseline["combined_clv"]
        print(f"    {vs['label']} Period Totals:")
        print(f"      Combined Net Revenue:       {_fd(vs['combined_net'])}")
        print(f"      Combined CLV Lost:           {_fd(vs['combined_clv'])}")
        print(f"      Combined Gross Revenue:     {_fd(vs['combined_gross'])}")
        print(f"      Combined Gateway Cost:      {_fd(vs['combined_gw_cost'])}")
        print(f"      Combined Refunds:           {_fd(vs['combined_refunds'])}")
        print(f"      Combined Chargebacks:       {_fd(vs['combined_chargebacks'])}")
        if vs is not baseline:
            print(f"      Delta Net vs Baseline:      {_fd(delta_net)}  {'(gain)' if delta_net > 0 else '(loss)'}")
            print(f"      Delta CLV Lost vs Baseline:  {_fd(delta_clv)}  {'(worse)' if delta_clv > 0 else '(better)'}")
        print()

    # Summary table
    print(f"\n  {'=' * 90}")
    print(f"  VOLUME-SHIFT SUMMARY ({len(monthly)}-Month Period)")
    print(f"  {'=' * 90}")
    print(f"  {'Scenario':<30} {'Net Revenue':>18} {'CLV Lost':>18} {'Net - CLV Lost':>18}")
    print(f"  {'-' * 90}")
    for vs in vs_results:
        net_minus_clv = vs["combined_net"] - vs["combined_clv"]
        print(f"  {vs['label']:<30} {_fd(vs['combined_net']):>18} {_fd(vs['combined_clv']):>18} {_fd(net_minus_clv):>18}")

    print(f"\n  {'Scenario':<30} {'Delta Net':>18} {'Delta CLV Lost':>18} {'Delta Net-CLV':>18}")
    print(f"  {'-' * 90}")
    for vs in vs_results[1:]:
        dn = vs["combined_net"] - baseline["combined_net"]
        dc = vs["combined_clv"] - baseline["combined_clv"]
        dnc = dn - dc
        print(f"  {vs['label']:<30} {_fd(dn):>18} {_fd(dc):>18} {_fd(dnc):>18}")


# ═══════════════════════════════════════════════════════════════════════════
#  EXCEL OUTPUT
# ═══════════════════════════════════════════════════════════════════════════

def write_excel(monthly, totals, wi_months, wi_totals, vs_results, params, path):
    if not HAS_OPENPYXL:
        print("\n[WARN] openpyxl not installed -- skipping Excel output.")
        return

    wb = Workbook()
    hdr_font = Font(bold=True, size=11)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    money_fmt = '#,##0.00'
    pct_fmt = '0.00%'
    int_fmt = '#,##0'
    thin = Side(style="thin")
    border = Border(bottom=thin)

    def style_header(ws, row_num, max_col):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # ── Sheet 1: Monthly Comparison ──
    ws = wb.active
    ws.title = "Monthly Comparison"
    headers = [
        "Month",
        f"DTP Attempts ({GW_A})", f"DTP Attempts ({GW_C})",
        f"DTP Success Rate ({GW_A})", f"DTP Success Rate ({GW_C})",
        f"Ren Attempts ({GW_A})", f"Ren Attempts ({GW_C})",
        f"Ren Success Rate ({GW_A})", f"Ren Success Rate ({GW_C})",
        f"Gross Revenue ({GW_A})", f"Gross Revenue ({GW_C})",
        f"Gateway Cost ({GW_A})", f"Gateway Cost ({GW_C})",
        f"Refunds ({GW_A})", f"Refunds ({GW_C})",
        f"Chargebacks ({GW_A})", f"Chargebacks ({GW_C})",
        f"Net Revenue ({GW_A})", f"Net Revenue ({GW_C})",
        f"Net/Attempt ({GW_A})", f"Net/Attempt ({GW_C})",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(headers))

    all_rows = monthly + [totals]
    for ri, r in enumerate(all_rows, 2):
        a, c = r["a"], r["c"]
        vals = [
            r["label"],
            a["vol_ci"], c["vol_ci"],
            a["sr_ci"], c["sr_ci"],
            a["vol_ren"], c["vol_ren"],
            a["sr_ren"], c["sr_ren"],
            a["gross"], c["gross"],
            a["gw_cost"], c["gw_cost"],
            a["refunds"], c["refunds"],
            a["chargebacks"], c["chargebacks"],
            a["net"], c["net"],
            a["net_per_attempt"], c["net_per_attempt"],
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            if ci in (4, 5, 8, 9):
                cell.number_format = pct_fmt
            elif ci in (2, 3, 6, 7):
                cell.number_format = int_fmt
            elif ci >= 10:
                cell.number_format = money_fmt
        if r is all_rows[-1]:
            for ci in range(1, len(headers) + 1):
                ws.cell(row=ri, column=ci).font = hdr_font

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    # ── Sheet 2: CLV Impact ──
    ws2 = wb.create_sheet("CLV Lost to Failures")
    clv_headers = [
        "Month",
        f"Failed DTP ({GW_A})", f"Failed DTP ({GW_C})",
        f"DTP Immediate Loss ({GW_A})", f"DTP Immediate Loss ({GW_C})",
        f"DTP Future Loss ({GW_A})", f"DTP Future Loss ({GW_C})",
        f"Failed Ren ({GW_A})", f"Failed Ren ({GW_C})",
        f"Ren Loss ({GW_A})", f"Ren Loss ({GW_C})",
        f"Total CLV Lost ({GW_A})", f"Total CLV Lost ({GW_C})",
    ]
    for i, h in enumerate(clv_headers, 1):
        ws2.cell(row=1, column=i, value=h)
    style_header(ws2, 1, len(clv_headers))

    for ri, r in enumerate(all_rows, 2):
        a, c = r["a"], r["c"]
        vals = [
            r["label"],
            a["fail_ci"], c["fail_ci"],
            a["lost_imm_ci"], c["lost_imm_ci"],
            a["lost_fut_ci"], c["lost_fut_ci"],
            a["fail_ren"], c["fail_ren"],
            a["lost_ren"], c["lost_ren"],
            a["clv_impact"], c["clv_impact"],
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws2.cell(row=ri, column=ci, value=v)
            if ci in (2, 3, 8, 9):
                cell.number_format = int_fmt
            elif ci >= 4:
                cell.number_format = money_fmt

    for col in range(1, len(clv_headers) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 22

    # ── Sheet 3: What-If ──
    ws3 = wb.create_sheet("What-If Analysis")
    wi_headers = [
        "Month", "Combined DTP Vol", "Combined Ren Vol",
        f"Net Revenue (All {GW_A})", f"Net Revenue (All {GW_C})", "Net Delta (A-C)",
        f"CLV Lost (All {GW_A})", f"CLV Lost (All {GW_C})", "CLV Lost Delta (A-C)",
    ]
    for i, h in enumerate(wi_headers, 1):
        ws3.cell(row=1, column=i, value=h)
    style_header(ws3, 1, len(wi_headers))

    all_wi = wi_months + [wi_totals]
    for ri, w in enumerate(all_wi, 2):
        a, c = w["a"], w["c"]
        vals = [
            w["label"], w["tot_ci"], w["tot_ren"],
            a["net"], c["net"], a["net"] - c["net"],
            a["clv"], c["clv"], a["clv"] - c["clv"],
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws3.cell(row=ri, column=ci, value=v)
            if ci in (2, 3):
                cell.number_format = int_fmt
            elif ci >= 4:
                cell.number_format = money_fmt
        if w is all_wi[-1]:
            for ci in range(1, len(wi_headers) + 1):
                ws3.cell(row=ri, column=ci).font = hdr_font

    for col in range(1, len(wi_headers) + 1):
        ws3.column_dimensions[get_column_letter(col)].width = 24

    # ── Sheet 4: Volume-Shift Scenarios ──
    ws4 = wb.create_sheet("Volume-Shift Scenarios")
    vs_headers = [
        "Scenario",
        "Combined Net Revenue", "Combined CLV Lost",
        "Combined Gross Revenue", "Combined Gateway Cost",
        "Combined Refunds", "Combined Chargebacks",
        f"{GW_A} Net Revenue", f"{GW_C} Net Revenue",
        "Delta Net vs Baseline", "Delta CLV Lost vs Baseline",
    ]
    for i, h in enumerate(vs_headers, 1):
        ws4.cell(row=1, column=i, value=h)
    style_header(ws4, 1, len(vs_headers))

    baseline = vs_results[0]
    for ri, vs in enumerate(vs_results, 2):
        dn = vs["combined_net"] - baseline["combined_net"]
        dc = vs["combined_clv"] - baseline["combined_clv"]
        vals = [
            vs["label"],
            vs["combined_net"], vs["combined_clv"],
            vs["combined_gross"], vs["combined_gw_cost"],
            vs["combined_refunds"], vs["combined_chargebacks"],
            vs["adyen_net"], vs["cyber_net"],
            dn, dc,
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws4.cell(row=ri, column=ci, value=v)
            if ci >= 2:
                cell.number_format = money_fmt

    for col in range(1, len(vs_headers) + 1):
        ws4.column_dimensions[get_column_letter(col)].width = 24

    # ── Sheet 5: Parameters ──
    ws5 = wb.create_sheet("Parameters")
    param_data = [
        ("Parameter", "Value"),
        (f"{GW_A} Cost Rate", params["cost_a"]),
        (f"{GW_C} Cost Rate", params["cost_c"]),
        ("Average Transaction Amount", params["avg_txn"]),
        ("Customer Lifetime (months)", params["lifetime"]),
        ("CLV per Customer", params["lifetime"] * params["avg_txn"]),
        ("Future Renewals Lost per Failed DTP", (params["lifetime"] - 1) * params["avg_txn"]),
    ]
    for ri, (label, val) in enumerate(param_data, 1):
        ws5.cell(row=ri, column=1, value=label)
        cell = ws5.cell(row=ri, column=2, value=val)
        if ri == 1:
            ws5.cell(row=ri, column=1).font = hdr_font
            cell.font = hdr_font
        elif ri in (2, 3):
            cell.number_format = pct_fmt
        elif ri >= 4:
            cell.number_format = money_fmt
    ws5.column_dimensions["A"].width = 35
    ws5.column_dimensions["B"].width = 18

    wb.save(path)
    print(f"\n[OK] Excel report saved to {path}")


# ═══════════════════════════════════════════════════════════════════════════
#  CHART GENERATION
# ═══════════════════════════════════════════════════════════════════════════

def generate_charts(monthly, totals, wi_months, wi_totals, vs_results, params, out_dir):
    if not HAS_MPL:
        print("\n[WARN] matplotlib not installed -- skipping chart generation.")
        return

    os.makedirs(out_dir, exist_ok=True)
    plt.style.use("seaborn-v0_8-whitegrid")

    labels = [r["label"] for r in monthly]
    x = np.arange(len(labels))

    _chart_success_rates(monthly, labels, x, out_dir)
    _chart_net_revenue(monthly, labels, x, out_dir)
    _chart_clv_impact(monthly, labels, x, out_dir)
    _chart_cost_waterfall(totals, params, out_dir)
    _chart_what_if(totals, wi_totals, params, out_dir)
    _chart_volume_shift(vs_results, totals, params, out_dir)

    print(f"[OK] Charts saved to {out_dir}/")


def _chart_success_rates(monthly, labels, x, out_dir):
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(FIG_W, FIG_H))
    w = 0.35

    ci_a = [r["a"]["sr_ci"] * 100 for r in monthly]
    ci_c = [r["c"]["sr_ci"] * 100 for r in monthly]
    ax1.bar(x - w / 2, ci_a, w, label=GW_A, color=CLR_ADYEN)
    ax1.bar(x + w / 2, ci_c, w, label=GW_C, color=CLR_CYBER)
    for i in range(len(x)):
        ax1.text(x[i] - w / 2, ci_a[i] + 0.5, f"{ci_a[i]:.1f}%", ha="center", fontsize=8)
        ax1.text(x[i] + w / 2, ci_c[i] + 0.5, f"{ci_c[i]:.1f}%", ha="center", fontsize=8)
    ax1.set_xticks(x)
    ax1.set_xticklabels(labels)
    ax1.set_ylabel("Success Rate (%)")
    ax1.set_title("Direct-to-Paid Success Rates")
    ax1.legend()
    ax1.set_ylim(0, 105)

    ren_a = [r["a"]["sr_ren"] * 100 for r in monthly]
    ren_c = [r["c"]["sr_ren"] * 100 for r in monthly]
    ax2.bar(x - w / 2, ren_a, w, label=GW_A, color=CLR_ADYEN)
    ax2.bar(x + w / 2, ren_c, w, label=GW_C, color=CLR_CYBER)
    for i in range(len(x)):
        ax2.text(x[i] - w / 2, ren_a[i] + 0.3, f"{ren_a[i]:.1f}%", ha="center", fontsize=8)
        ax2.text(x[i] + w / 2, ren_c[i] + 0.3, f"{ren_c[i]:.1f}%", ha="center", fontsize=8)
    ax2.set_xticks(x)
    ax2.set_xticklabels(labels)
    ax2.set_ylabel("Success Rate (%)")
    ax2.set_title("Renewal Success Rates")
    ax2.legend()
    ax2.set_ylim(85, 100)

    fig.suptitle("Monthly Success Rates by Gateway", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    fig.savefig(os.path.join(out_dir, "1_success_rates.png"), dpi=DPI)
    plt.close(fig)


def _chart_net_revenue(monthly, labels, x, out_dir):
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    w = 0.35

    net_a = [r["a"]["net"] for r in monthly]
    net_c = [r["c"]["net"] for r in monthly]
    cum_a = np.cumsum(net_a)
    cum_c = np.cumsum(net_c)

    ax.bar(x - w / 2, net_a, w, label=f"{GW_A} (monthly)", color=CLR_ADYEN)
    ax.bar(x + w / 2, net_c, w, label=f"{GW_C} (monthly)", color=CLR_CYBER)

    ax2 = ax.twinx()
    ax2.plot(x, cum_a, "o--", color=CLR_ADYEN, label=f"{GW_A} (cumulative)", linewidth=2)
    ax2.plot(x, cum_c, "s--", color=CLR_CYBER, label=f"{GW_C} (cumulative)", linewidth=2)

    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("Monthly Net Revenue ($)")
    ax2.set_ylabel("Cumulative Net Revenue ($)")
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))
    ax2.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    lines1, labels1 = ax.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax.legend(lines1 + lines2, labels1 + labels2, loc="upper left")

    ax.set_title("Monthly Net Revenue Comparison\n(after gateway costs, refunds, and chargebacks)",
                 fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(os.path.join(out_dir, "2_net_revenue.png"), dpi=DPI)
    plt.close(fig)


def _chart_clv_impact(monthly, labels, x, out_dir):
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    w = 0.35

    imm_a = [r["a"]["lost_imm_ci"] for r in monthly]
    fut_a = [r["a"]["lost_fut_ci"] for r in monthly]
    ren_a = [r["a"]["lost_ren"] for r in monthly]

    imm_c = [r["c"]["lost_imm_ci"] for r in monthly]
    fut_c = [r["c"]["lost_fut_ci"] for r in monthly]
    ren_c = [r["c"]["lost_ren"] for r in monthly]

    ax.bar(x - w / 2, imm_a, w, label=f"DTP Immediate ({GW_A})", color=CLR_ADYEN)
    ax.bar(x - w / 2, fut_a, w, bottom=imm_a, label=f"DTP Future Renewals ({GW_A})", color=CLR_ADYEN_LT)
    ax.bar(x - w / 2, ren_a, w,
           bottom=[a + b for a, b in zip(imm_a, fut_a)],
           label=f"Renewal Loss ({GW_A})", color=CLR_GRAY, alpha=0.5)

    ax.bar(x + w / 2, imm_c, w, label=f"DTP Immediate ({GW_C})", color=CLR_CYBER)
    ax.bar(x + w / 2, fut_c, w, bottom=imm_c, label=f"DTP Future Renewals ({GW_C})", color=CLR_CYBER_LT)
    ax.bar(x + w / 2, ren_c, w,
           bottom=[a + b for a, b in zip(imm_c, fut_c)],
           label=f"Renewal Loss ({GW_C})", color=CLR_GRAY, alpha=0.3)

    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("Revenue Impact ($)")
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))
    ax.legend(loc="upper right", fontsize=9)
    ax.set_title("CLV Lost to Transaction Failures\n(Failed DTP = lost customer x full lifetime value)",
                 fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(os.path.join(out_dir, "3_clv_impact.png"), dpi=DPI)
    plt.close(fig)


def _chart_cost_waterfall(totals, params, out_dir):
    fig, (ax_a, ax_c) = plt.subplots(1, 2, figsize=(FIG_W, FIG_H))

    for ax, gw, color, name in [
        (ax_a, "a", CLR_ADYEN, f"{GW_A} ({params['cost_a']:.1%})"),
        (ax_c, "c", CLR_CYBER, f"{GW_C} ({params['cost_c']:.1%})"),
    ]:
        t = totals[gw]
        cats = ["Gross\nRevenue", "Gateway\nCost", "Refunds", "Chargebacks", "Net\nRevenue"]
        values = [t["gross"], -t["gw_cost"], -t["refunds"], -t["chargebacks"], t["net"]]
        bottoms = [0, t["gross"] - t["gw_cost"], t["gross"] - t["gw_cost"] - t["refunds"],
                   t["gross"] - t["gw_cost"] - t["refunds"] - t["chargebacks"], 0]
        colors = [CLR_GREEN, CLR_RED, CLR_RED, CLR_RED, color]
        heights = [t["gross"], t["gw_cost"], t["refunds"], t["chargebacks"], t["net"]]

        bars = ax.bar(cats, heights, bottom=bottoms, color=colors, alpha=0.8,
                      edgecolor="white", linewidth=1.5)

        for bar, val in zip(bars, heights):
            y = bar.get_y() + bar.get_height() / 2
            ax.text(bar.get_x() + bar.get_width() / 2, y,
                    f"${val:,.0f}", ha="center", va="center", fontsize=9,
                    fontweight="bold", color="white")

        connector_tops = [
            t["gross"],
            t["gross"] - t["gw_cost"],
            t["gross"] - t["gw_cost"] - t["refunds"],
        ]
        for ci in range(len(connector_tops)):
            ax.plot([ci + 0.4, ci + 0.6], [connector_tops[ci], connector_tops[ci]],
                    "k-", linewidth=0.8, alpha=0.4)

        ax.set_title(name, fontsize=12, fontweight="bold")
        ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    fig.suptitle(f"{totals['label']} Cost Breakdown Waterfall", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    fig.savefig(os.path.join(out_dir, "4_cost_waterfall.png"), dpi=DPI)
    plt.close(fig)


def _chart_what_if(totals, wi_totals, params, out_dir):
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(FIG_W, FIG_H))

    actual_net = totals["a"]["net"] + totals["c"]["net"]
    actual_clv = totals["a"]["clv_impact"] + totals["c"]["clv_impact"]
    all_a_net = wi_totals["a"]["net"]
    all_c_net = wi_totals["c"]["net"]
    all_a_clv = wi_totals["a"]["clv"]
    all_c_clv = wi_totals["c"]["clv"]

    scenarios = ["Actual\n(Split)", f"All via {GW_A}\n({params['cost_a']:.1%})", f"All via {GW_C}\n({params['cost_c']:.1%})"]
    nets = [actual_net, all_a_net, all_c_net]
    clvs = [actual_clv, all_a_clv, all_c_clv]
    colors = [CLR_GRAY, CLR_ADYEN, CLR_CYBER]

    bars1 = ax1.barh(scenarios, nets, color=colors, alpha=0.85)
    for bar, val in zip(bars1, nets):
        ax1.text(bar.get_width() * 0.98, bar.get_y() + bar.get_height() / 2,
                 f"${val:,.0f}", ha="right", va="center", fontsize=10, fontweight="bold", color="white")
    ax1.set_xlabel("Net Revenue ($)")
    ax1.set_title("Total Net Revenue", fontsize=12, fontweight="bold")
    ax1.xaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    bars2 = ax2.barh(scenarios, clvs, color=colors, alpha=0.85)
    for bar, val in zip(bars2, clvs):
        ax2.text(bar.get_width() * 0.98, bar.get_y() + bar.get_height() / 2,
                 f"${val:,.0f}", ha="right", va="center", fontsize=10, fontweight="bold", color="white")
    ax2.set_xlabel("CLV Lost to Failures ($)")
    ax2.set_title("Total CLV Lost to Failures (lower = better)", fontsize=12, fontweight="bold")
    ax2.xaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    fig.suptitle(f"{totals['label']} What-If Scenario Comparison", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    fig.savefig(os.path.join(out_dir, "5_what_if_scenarios.png"), dpi=DPI)
    plt.close(fig)


def _chart_volume_shift(vs_results, totals, params, out_dir):
    fig, axes = plt.subplots(1, 3, figsize=(FIG_W + 4, FIG_H))

    scenario_labels = [vs["label"] for vs in vs_results]
    scenario_short = [f"{vs['pct_a']:.0%}" for vs in vs_results]
    nets = [vs["combined_net"] for vs in vs_results]
    clvs = [vs["combined_clv"] for vs in vs_results]
    net_minus_clv = [n - c for n, c in zip(nets, clvs)]

    colors = [CLR_GRAY, CLR_ADYEN, CLR_PURPLE]

    # Net Revenue
    bars1 = axes[0].bar(scenario_short, nets, color=colors, alpha=0.85, edgecolor="white")
    for bar, val in zip(bars1, nets):
        axes[0].text(bar.get_x() + bar.get_width() / 2, bar.get_height(),
                     f"${val:,.0f}", ha="center", va="bottom", fontsize=8, fontweight="bold")
    axes[0].set_title("Combined Net Revenue", fontsize=11, fontweight="bold")
    axes[0].set_xlabel(f"{GW_A} DTP Share")
    axes[0].yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    # CLV Lost to Failures
    bars2 = axes[1].bar(scenario_short, clvs, color=colors, alpha=0.85, edgecolor="white")
    for bar, val in zip(bars2, clvs):
        axes[1].text(bar.get_x() + bar.get_width() / 2, bar.get_height(),
                     f"${val:,.0f}", ha="center", va="bottom", fontsize=8, fontweight="bold")
    axes[1].set_title("Combined CLV Lost to Failures\n(lower = better)", fontsize=11, fontweight="bold")
    axes[1].set_xlabel(f"{GW_A} DTP Share")
    axes[1].yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    # Net Revenue minus CLV
    bars3 = axes[2].bar(scenario_short, net_minus_clv, color=colors, alpha=0.85, edgecolor="white")
    for bar, val in zip(bars3, net_minus_clv):
        axes[2].text(bar.get_x() + bar.get_width() / 2, bar.get_height(),
                     f"${val:,.0f}", ha="center", va="bottom", fontsize=8, fontweight="bold")
    axes[2].set_title("Net Revenue minus CLV Lost", fontsize=11, fontweight="bold")
    axes[2].set_xlabel(f"{GW_A} DTP Share")
    axes[2].yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))

    baseline = vs_results[0]
    for vs in vs_results[1:]:
        delta_net = vs["combined_net"] - baseline["combined_net"]
        delta_clv = vs["combined_clv"] - baseline["combined_clv"]
        delta_net_clv = delta_net - delta_clv

    fig.suptitle(f"Volume-Shift Scenario Impact ({totals['label']})\n"
                 f"Shifting {GW_A} DTP share from 10% to 20% or 25%",
                 fontsize=14, fontweight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.88])
    fig.savefig(os.path.join(out_dir, "6_volume_shift.png"), dpi=DPI)
    plt.close(fig)


# ═══════════════════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(description="Gateway Performance Comparison Tool")
    parser.add_argument("input_file", help="Path to the gateway performance xlsx file")
    parser.add_argument("--output-dir", "-o", default=None,
                        help="Output directory (default: gateway_comparison_output/ next to input)")
    args = parser.parse_args()

    input_path = Path(args.input_file).resolve()
    if not input_path.exists():
        print(f"Error: {input_path} not found")
        sys.exit(1)

    out_dir = Path(args.output_dir) if args.output_dir else input_path.parent / "gateway_comparison_output"
    os.makedirs(out_dir, exist_ok=True)

    print(f"Loading data from {input_path} ...")
    params, month_list = load_data(str(input_path))
    print(f"  {GW_A} cost: {params['cost_a']:.1%}")
    print(f"  {GW_C} cost: {params['cost_c']:.1%}")
    print(f"  Avg transaction: ${params['avg_txn']:.2f}")
    print(f"  Lifetime: {params['lifetime']} months")
    print(f"  Months loaded: {len(month_list)} ({month_list[0]['label']} -- {month_list[-1]['label']})")

    monthly = calculate_all(month_list, params)
    totals = aggregate_totals(monthly)
    wi_months, wi_totals = what_if(month_list, params)
    vs_results = volume_shift(month_list, params)

    print_report(monthly, totals, wi_months, wi_totals, vs_results, params)

    excel_path = out_dir / "gateway_comparison_results.xlsx"
    write_excel(monthly, totals, wi_months, wi_totals, vs_results, params, str(excel_path))

    chart_dir = str(out_dir / "charts")
    generate_charts(monthly, totals, wi_months, wi_totals, vs_results, params, chart_dir)


if __name__ == "__main__":
    main()
